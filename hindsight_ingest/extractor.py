from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


def extract_xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                lines.append("\t".join("" if cell is None else str(cell) for cell in row))
    wb.close()
    return "\n".join(lines)


def extract_pdf(path: Path) -> Optional[str]:
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(path))
        text = "\n".join(page.get_text() for page in doc)
        doc.close()
        return text.strip() or None
    except ImportError:
        logger.warning("PyMuPDF not installed; cannot extract text from %s", path.name)
        return None
    except Exception as exc:
        logger.warning("PDF extraction failed for %s: %s", path.name, exc)
        return None


def extract_docx(path: Path) -> Optional[str]:
    try:
        from docx import Document
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text)
        return text.strip() or None
    except ImportError:
        logger.warning("python-docx not installed; cannot extract text from %s", path.name)
        return None
    except Exception as exc:
        logger.warning("DOCX extraction failed for %s: %s", path.name, exc)
        return None


def extract_pptx(path: Path) -> Optional[str]:
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        lines: list[str] = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(shape.text.strip())
        return "\n".join(lines) or None
    except ImportError:
        logger.warning("python-pptx not installed; cannot extract text from %s", path.name)
        return None
    except Exception as exc:
        logger.warning("PPTX extraction failed for %s: %s", path.name, exc)
        return None


def extract_image_with_tesseract(path: Path) -> Optional[str]:
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(str(path))
        text = pytesseract.image_to_string(img)
        return text.strip() or None
    except ImportError:
        logger.warning("pytesseract/Pillow not installed; cannot OCR %s", path.name)
        return None
    except Exception as exc:
        logger.warning("Tesseract OCR failed for %s: %s", path.name, exc)
        return None
