from pathlib import Path

import pytest

from hindsight_ingest.extractor import extract_xlsx, extract_image_with_tesseract
from hindsight_ingest.chunker import chunk_text


# ── XLSX ──────────────────────────────────────────────────────────────────────

def _make_xlsx(tmp_path, rows: list) -> Path:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TestSheet"
    for row in rows:
        ws.append(row)
    path = tmp_path / "test.xlsx"
    wb.save(str(path))
    return path


def test_extract_xlsx_sheet_name_present(tmp_path):
    path = _make_xlsx(tmp_path, [["Name", "Age"], ["Alice", 30]])
    text = extract_xlsx(path)
    assert "TestSheet" in text


def test_extract_xlsx_data_present(tmp_path):
    path = _make_xlsx(tmp_path, [["Name", "Score"], ["Bob", 99]])
    text = extract_xlsx(path)
    assert "Bob" in text
    assert "99" in text


def test_extract_xlsx_empty_rows_skipped(tmp_path):
    path = _make_xlsx(tmp_path, [["Data"], [None, None], ["More"]])
    text = extract_xlsx(path)
    lines = [ln for ln in text.splitlines() if ln.strip() == ""]
    assert len(lines) == 0


def test_extract_xlsx_multiple_sheets(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["hello"])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["world"])
    path = tmp_path / "multi.xlsx"
    wb.save(str(path))
    text = extract_xlsx(path)
    assert "Sheet1" in text and "Sheet2" in text
    assert "hello" in text and "world" in text


# ── Chunker ───────────────────────────────────────────────────────────────────

def test_chunk_text_basic():
    text = " ".join(str(i) for i in range(10))
    chunks = chunk_text(text, chunk_size=4, overlap=1)
    assert all(len(c.split()) <= 4 for c in chunks)
    assert chunks[0].startswith("0")


def test_chunk_text_empty_returns_empty():
    assert chunk_text("", 100, 10) == []


def test_chunk_text_overlap():
    words = list("abcdefghij")
    text = " ".join(words)
    chunks = chunk_text(text, chunk_size=4, overlap=2)
    # second chunk should start with 'c' (overlap of 2 from 'a b c d')
    assert chunks[1].split()[0] == "c"


def test_chunk_text_single_chunk_when_small():
    text = "one two three"
    chunks = chunk_text(text, chunk_size=100, overlap=10)
    assert len(chunks) == 1
    assert chunks[0] == text


# ── Tesseract (graceful degradation) ─────────────────────────────────────────

def test_tesseract_returns_none_if_not_installed(tmp_path, monkeypatch):
    monkeypatch.setitem(__import__("sys").modules, "pytesseract", None)
    from importlib import reload
    import hindsight_ingest.extractor as ext_mod
    # Simulate ImportError path by calling directly
    result = ext_mod.extract_image_with_tesseract(tmp_path / "ghost.png")
    assert result is None
